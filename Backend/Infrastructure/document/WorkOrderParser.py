import csv
import io
import json
import logging
import os
import uuid
from datetime import datetime

from dateutil import parser as date_parser

from Backend.Application.Interfaces.ILLMService import ILLMService
from Backend.Domain.Entities.work_order import WorkOrder

logger = logging.getLogger(__name__)

# 支持中英文表头与常见别名，减少结构化工单导入时的列名耦合。
_FIELD_ALIASES = {
    "order_no": {"order_no", "orderno", "orderid", "工单编号", "工单号", "编号", "单号"},
    "device_name": {"device_name", "devicename", "设备名称", "设备名", "设备"},
    "device_code": {"device_code", "devicecode", "设备编码", "设备代码", "设备编号"},
    "fault_phenomenon": {"fault_phenomenon", "faultphenomenon", "symptom", "故障现象", "故障描述", "现象"},
    "fault_cause": {"fault_cause", "faultcause", "rootcause", "故障原因", "原因", "根因"},
    "fault_category": {"fault_category", "faultcategory", "category", "故障分类", "分类"},
    "severity": {"severity", "level", "严重等级", "严重程度", "等级"},
    "solution": {"solution", "measure", "action", "处置措施", "处理措施", "解决方案", "措施"},
    "occurrence_time": {"occurrence_time", "occurrencetime", "发生时间", "故障时间", "开始时间"},
    "resolution_time": {"resolution_time", "resolutiontime", "解决时间", "恢复时间", "结束时间"},
    "operator": {"operator", "owner", "handler", "处理人", "维修人", "责任人"},
    "raw_text": {"raw_text", "rawtext", "原始文本", "原文", "全文"},
}

_TEXT_PARSE_SYSTEM = "你是工业设备工单结构化助手，只返回 JSON 对象，不要输出解释。"

_TEXT_PARSE_USER = """请从下面的工单文本中提取结构化字段，并严格返回 JSON 对象。

字段要求：
- order_no: 工单编号，没有则留空字符串
- device_name: 设备名称，没有则留空字符串
- device_code: 设备编码，没有则留空字符串
- fault_phenomenon: 故障现象描述，没有则留空字符串
- fault_cause: 故障原因，没有则留空字符串
- fault_category: 故障分类（机械/电气/液压/控制/软件），没有则留空字符串
- severity: 严重等级（一般/严重/紧急），没有则留空字符串
- solution: 处置措施，没有则留空字符串
- occurrence_time: ISO8601 时间字符串，没有则返回 null
- resolution_time: ISO8601 时间字符串，没有则返回 null
- operator: 处理人，没有则留空字符串
- raw_text: 原始文本原样返回

只返回如下格式的 JSON：
{{
  "order_no": "",
  "device_name": "",
  "device_code": "",
  "fault_phenomenon": "",
  "fault_cause": "",
  "fault_category": "",
  "severity": "",
  "solution": "",
  "occurrence_time": null,
  "resolution_time": null,
  "operator": "",
  "raw_text": ""
}}

工单文本：
{text}
"""


class WorkOrderParser:
    """工单导入解析器。

    统一处理三种输入：
    1. CSV 结构化表格。
    2. Excel 结构化表格。
    3. 自由文本工单，交给 LLM 抽取字段。
    """

    def __init__(self, llm_service: ILLMService):
        self._llm = llm_service
        self._alias_to_field = {
            self._normalize_header(alias): field
            for field, aliases in _FIELD_ALIASES.items()
            for alias in aliases
        }

    def parse_upload(self, uploaded_file, filename: str) -> tuple[list[WorkOrder], list[dict]]:
        """根据上传文件扩展名分派对应解析逻辑。"""
        ext = os.path.splitext(filename)[1].lower()
        payload = uploaded_file.read()
        if ext == ".csv":
            return self._parse_csv(payload, filename)
        if ext in {".xlsx", ".xls"}:
            return self._parse_excel(payload, filename, ext)
        if ext in {".txt", ".md"}:
            text = self._decode_bytes(payload)
            return [self.parse_text(text, filename)], []
        raise ValueError(f"Unsupported work order import type: {ext}")

    def parse_text(self, text: str, source_file: str = "manual_text.txt") -> WorkOrder:
        """调用 LLM 解析自由文本工单，并对缺失关键字段做兜底。"""
        messages = [
            {"role": "system", "content": _TEXT_PARSE_SYSTEM},
            {"role": "user", "content": _TEXT_PARSE_USER.format(text=text)},
        ]
        raw = ""
        for token in self._llm.stream_chat(messages):
            raw += token
        data = self._parse_json(raw)
        if not isinstance(data, dict):
            raise ValueError("Failed to parse work order text into JSON")
        data["raw_text"] = data.get("raw_text") or text
        data["source_file"] = source_file
        if not data.get("order_no"):
            data["order_no"] = f"AUTO-{uuid.uuid4().hex[:10].upper()}"
        if not data.get("device_name"):
            data["device_name"] = "未知设备"
        return self._build_work_order(data, source_file=source_file)

    def _parse_csv(self, payload: bytes, source_file: str) -> tuple[list[WorkOrder], list[dict]]:
        """解析 CSV 工单文件。"""
        text = self._decode_bytes(payload)
        reader = csv.DictReader(io.StringIO(text))
        return self._parse_rows(list(reader), source_file=source_file)

    def _parse_excel(
        self,
        payload: bytes,
        source_file: str,
        ext: str,
    ) -> tuple[list[WorkOrder], list[dict]]:
        """解析 Excel 工单文件。

        .xlsx 使用 openpyxl，.xls 使用 xlrd，避免单一库兼容性不足。
        """
        if ext == ".xlsx":
            try:
                from openpyxl import load_workbook
            except ImportError as exc:
                raise ImportError("openpyxl is required for .xlsx work order import") from exc
            workbook = load_workbook(io.BytesIO(payload), read_only=True, data_only=True)
            sheet = workbook.active
            rows = list(sheet.iter_rows(values_only=True))
            workbook.close()
            if not rows:
                return [], []
            headers = [str(cell).strip() if cell is not None else "" for cell in rows[0]]
            parsed_rows = []
            for row in rows[1:]:
                parsed_rows.append({
                    headers[index]: "" if value is None else str(value)
                    for index, value in enumerate(row)
                    if index < len(headers)
                })
            return self._parse_rows(parsed_rows, source_file=source_file)

        try:
            import xlrd
        except ImportError as exc:
            raise ImportError("xlrd is required for .xls work order import") from exc
        workbook = xlrd.open_workbook(file_contents=payload)
        sheet = workbook.sheet_by_index(0)
        if sheet.nrows == 0:
            return [], []
        headers = [str(sheet.cell_value(0, col)).strip() for col in range(sheet.ncols)]
        parsed_rows = []
        for row_idx in range(1, sheet.nrows):
            parsed_rows.append({
                headers[col]: "" if sheet.cell_value(row_idx, col) is None else str(sheet.cell_value(row_idx, col))
                for col in range(sheet.ncols)
            })
        return self._parse_rows(parsed_rows, source_file=source_file)

    def _parse_rows(self, rows: list[dict], source_file: str) -> tuple[list[WorkOrder], list[dict]]:
        """逐行解析结构化工单，并收集错误而不是中断整批导入。"""
        work_orders: list[WorkOrder] = []
        errors: list[dict] = []

        for index, row in enumerate(rows, start=2):
            try:
                normalized = self._normalize_row(row)
                # 结构化导入首版强制要求工单编号和设备名称，避免脏数据直接入库。
                if not normalized.get("order_no"):
                    raise ValueError("Missing required field: order_no")
                if not normalized.get("device_name"):
                    raise ValueError("Missing required field: device_name")
                work_orders.append(self._build_work_order(normalized, source_file=source_file))
            except Exception as exc:
                errors.append({"row": index, "error": str(exc)})

        return work_orders, errors

    def _normalize_row(self, row: dict) -> dict:
        """把原始表头映射为系统内部字段名。"""
        normalized: dict[str, str] = {}
        for raw_key, raw_value in row.items():
            field_name = self._alias_to_field.get(self._normalize_header(raw_key))
            if not field_name:
                continue
            normalized[field_name] = "" if raw_value is None else str(raw_value).strip()
        return normalized

    def _build_work_order(self, data: dict, source_file: str) -> WorkOrder:
        """把解析结果组装成 WorkOrder 实体。"""
        occurrence_time = self._parse_datetime(data.get("occurrence_time"))
        resolution_time = self._parse_datetime(data.get("resolution_time"))
        return WorkOrder(
            order_no=(data.get("order_no") or "").strip(),
            device_name=(data.get("device_name") or "").strip(),
            device_code=(data.get("device_code") or "").strip(),
            fault_phenomenon=(data.get("fault_phenomenon") or "").strip(),
            fault_cause=(data.get("fault_cause") or "").strip(),
            fault_category=(data.get("fault_category") or "").strip(),
            severity=(data.get("severity") or "").strip(),
            solution=(data.get("solution") or "").strip(),
            occurrence_time=occurrence_time,
            resolution_time=resolution_time,
            operator=(data.get("operator") or "").strip(),
            source_file=source_file,
            raw_text=(data.get("raw_text") or "").strip(),
        )

    @staticmethod
    def _normalize_header(header: str) -> str:
        """标准化表头，便于做别名匹配。"""
        if not header:
            return ""
        return (
            str(header)
            .strip()
            .lower()
            .replace(" ", "")
            .replace("_", "")
            .replace("-", "")
        )

    @staticmethod
    def _decode_bytes(payload: bytes) -> str:
        """按常见中文文件编码顺序尝试解码上传内容。"""
        for encoding in ("utf-8-sig", "utf-8", "gbk", "gb2312"):
            try:
                return payload.decode(encoding)
            except UnicodeDecodeError:
                continue
        return payload.decode("utf-8", errors="ignore")

    @staticmethod
    def _parse_datetime(value: str | None) -> datetime | None:
        """将字符串时间解析成 datetime；空值直接返回 None。"""
        if value in (None, "", "None"):
            return None
        if isinstance(value, datetime):
            return value
        text = str(value).strip()
        if not text:
            return None
        return date_parser.parse(text)

    @staticmethod
    def _parse_json(raw: str):
        """解析 LLM 返回的 JSON，并兼容 markdown 代码块包装。"""
        text = raw.strip()
        if text.startswith("```"):
            text = text.split("```", maxsplit=2)[1]
            if text.startswith("json"):
                text = text[4:]
        text = text.strip().rstrip("```").strip()
        try:
            return json.loads(text)
        except json.JSONDecodeError:
            logger.warning("WorkOrder text JSON parse failed: %s", raw[:200])
            return None
